import json
import os
import re
from io import BytesIO
from typing import Any

import pdfplumber
import google.generativeai as genai
from fastapi import APIRouter, Form, HTTPException, UploadFile, status
from pptx import Presentation

router = APIRouter(tags=["extract"])

# ---------------------------------------------------------------------------
# Gemini configuration
# ---------------------------------------------------------------------------

_EXTRACTION_PROMPT = """
You are analyzing an Anglepoint ROAR (Risk & Opportunity Assessment Report) or ELP (Enterprise License Position) document — an IT Asset Management consulting deliverable.
Extract all ROI metrics and return them as a single JSON object.

Return ONLY valid JSON, no markdown fences, no other text:
{
  "client": "client company name",
  "publisher": "software publisher (e.g. IBM, Microsoft, Oracle)",
  "date_delivered": "date as written in the document",
  "year": "4-digit year",
  "currency": "USD",
  "document_type": "ROAR or ELP",
  "identified_risk": 0,
  "identified_cost_avoidance": 0,
  "accomplished_cost_avoidance": 0,
  "identified_cost_optimization": 0,
  "accomplished_cost_optimization": 0,
  "realized_cost_savings": 0,
  "annual_publisher_contract_spend": 0,
  "pricing_available": "Yes",
  "notes": "",
  "elevate_deliverable": "Yes",
  "breakdown": [
    {
      "product": "product or topic name",
      "category": "Risk|Cost Avoidance|Cost Optimization|Savings",
      "identified": 0,
      "accomplished": 0,
      "description": "one-line description"
    }
  ]
}

Rules:
- All monetary values: plain numbers only, no $ or commas, use 0 if not found
- Empty string for missing text fields
- Populate breakdown with EVERY distinct financial line item found
- Return ONLY the JSON object, nothing else
"""


def _get_gemini_model() -> genai.GenerativeModel:
    """Configure and return a Gemini GenerativeModel, raising a helpful error if the key is missing."""
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=(
                "GEMINI_API_KEY environment variable is not set. "
                "Copy .env.example to .env and add your Gemini API key."
            ),
        )
    genai.configure(api_key=api_key)
    return genai.GenerativeModel("gemini-1.5-flash")


# ---------------------------------------------------------------------------
# Text extraction helpers
# ---------------------------------------------------------------------------

def _extract_text_from_pdf(content: bytes) -> str:
    """Extract all text from a PDF using pdfplumber."""
    pages: list[str] = []
    with pdfplumber.open(BytesIO(content)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(f"--- Page {i} ---\n{page_text}")
    return "\n\n".join(pages)


def _extract_text_from_pptx(content: bytes) -> str:
    """Extract all text from a PPTX using python-pptx."""
    prs = Presentation(BytesIO(content))
    slides: list[str] = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs).strip()
                    if line:
                        texts.append(line)
        if texts:
            slides.append(f"--- Slide {slide_num} ---\n" + "\n".join(texts))
    return "\n\n".join(slides)


# ---------------------------------------------------------------------------
# Gemini call
# ---------------------------------------------------------------------------

def _call_gemini(model: genai.GenerativeModel, document_text: str) -> dict[str, Any]:
    """Send document text to Gemini and return the parsed JSON dict."""
    prompt = document_text + "\n\n" + _EXTRACTION_PROMPT
    response = model.generate_content(prompt)
    raw = response.text.strip()

    # Strip markdown code fences if Gemini wraps the JSON anyway
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.IGNORECASE)
    raw = re.sub(r"\s*```$", "", raw)

    return json.loads(raw)


# ---------------------------------------------------------------------------
# Endpoint
# ---------------------------------------------------------------------------

@router.post("/")
async def extract_roi(
    file: UploadFile,
    document_type: str = Form(default="ROAR"),
    publisher: str = Form(default=""),
) -> dict[str, Any]:
    """
    Accept a PDF or PPTX file and return structured ROI data extracted by Gemini.
    """
    filename = file.filename or ""
    content = await file.read()

    # Determine file type from extension or MIME type
    lower_name = filename.lower()
    content_type = (file.content_type or "").lower()

    try:
        if lower_name.endswith(".pdf") or "pdf" in content_type:
            document_text = _extract_text_from_pdf(content)
        elif (
            lower_name.endswith(".pptx")
            or lower_name.endswith(".ppt")
            or "presentation" in content_type
            or "powerpoint" in content_type
        ):
            document_text = _extract_text_from_pptx(content)
        else:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Unsupported file type '{filename}'. Please upload a PDF or PPTX file.",
            )
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to extract text from file: {exc}",
        ) from exc

    if not document_text.strip():
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="No readable text found in the uploaded document.",
        )

    # Prepend any caller-supplied context so Gemini can use it
    context_lines: list[str] = []
    if document_type:
        context_lines.append(f"Document type: {document_type}")
    if publisher:
        context_lines.append(f"Publisher (if known): {publisher}")
    if context_lines:
        document_text = "\n".join(context_lines) + "\n\n" + document_text

    try:
        model = _get_gemini_model()
        result = _call_gemini(model, document_text)
    except HTTPException:
        raise
    except json.JSONDecodeError as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Gemini returned non-JSON output: {exc}",
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Gemini API error: {exc}",
        ) from exc

    return result
